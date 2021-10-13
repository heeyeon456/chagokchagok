import openpyxl
from openpyxl.chart import LineChart, Reference, BarChart
from openpyxl.drawing.image import Image
def xlWrite():
    wb = openpyxl.Workbook()
    sh = wb.active
    sh.append(['이름','국어','영어'])
    sh.append(['홍길동',40,80])
    sh.append(['이순신',50,90])
    sh.append(['임꺽정',70,60])
    wb.save('chart.xlsx')
#xlWrite()
wb = openpyxl.load_workbook('chart.xlsx')
sh = wb['Sheet']
#chart = LineChart()
chart = BarChart()
cat = Reference(sh, min_col=1, min_row=2, max_row=4)
cData = Reference(sh, min_col=2, max_col=3,
                  min_row=1, max_row=4)
chart.add_data(cData, titles_from_data=True)
chart.set_categories(cat)
sh.add_chart(chart, 'F1') # 대상 셀 위치
wb.save('chart.xlsx')

